"""导出层：把选中的 / 全部发票记录写成 CSV / XLSX / JSON。

- CSV 默认使用 utf-8-sig，Excel 直接双击打开不乱码；
- 金额列写入的是数值，方便在 Excel 里直接求和；
- 可选择是否附加"合计"行。
"""
from __future__ import annotations

import csv
import datetime as _dt
import json
import os
from typing import Any, Dict, List, Optional, Sequence, Tuple

import config
from core.models import InvoiceRecord


class ExportError(RuntimeError):
    pass


def _columns(keys: Optional[Sequence[str]]) -> List[Tuple[str, str]]:
    keys = list(keys) if keys else list(config.EXPORT_COLUMN_KEYS)
    pairs = [(k, label) for k, label in config.EXPORT_COLUMNS if k in keys]
    return pairs or list(config.EXPORT_COLUMNS)


def build_rows(
    records: Sequence[InvoiceRecord], keys: Optional[Sequence[str]] = None, formatted: bool = False
) -> Tuple[List[str], List[List[Any]]]:
    """返回 (表头, 数据行)。"""
    pairs = _columns(keys)
    headers = [label for _k, label in pairs]
    rows: List[List[Any]] = []
    for record in records:
        row = []
        for key, _label in pairs:
            if key in config.MONEY_FIELDS and not formatted:
                row.append(getattr(record, key, None))
            elif key == "confidence":
                row.append(round(float(record.confidence or 0), 3))
            else:
                row.append(record.value(key, formatted=formatted))
        rows.append(row)
    return headers, rows


def amount_sums(records: Sequence[InvoiceRecord]) -> Dict[str, float]:
    sums = {"amount": 0.0, "tax": 0.0, "total": 0.0}
    for record in records:
        for key in sums:
            value = getattr(record, key, None)
            if isinstance(value, (int, float)):
                sums[key] += float(value)
    return {k: round(v, 2) for k, v in sums.items()}


def _summary_line(headers: List[str], sums: Dict[str, float]) -> List[Any]:
    line: List[Any] = []
    for index, header in enumerate(headers):
        if index == 0:
            line.append("合计")
            continue
        key = None
        for k, label in config.EXPORT_COLUMNS:
            if label == header:
                key = k
                break
        line.append(sums.get(key) if key in sums else "")
    return line


# --------------------------------------------------------------------------
# 各格式写出
# --------------------------------------------------------------------------
def export_csv(
    path: str,
    records: Sequence[InvoiceRecord],
    keys: Optional[Sequence[str]] = None,
    encoding: str = "utf-8-sig",
    with_summary: bool = True,
) -> str:
    headers, rows = build_rows(records, keys)
    directory = os.path.dirname(os.path.abspath(path))
    if directory:
        os.makedirs(directory, exist_ok=True)
    with open(path, "w", newline="", encoding=encoding or "utf-8-sig", errors="replace") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
        if with_summary and records:
            sums = amount_sums(records)
            writer.writerow([])
            writer.writerow(_summary_line(headers, sums))
    return path


def export_xlsx(
    path: str,
    records: Sequence[InvoiceRecord],
    keys: Optional[Sequence[str]] = None,
    with_summary: bool = True,
) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise ExportError("导出 .xlsx 需要 openpyxl：pip install openpyxl（或改选 CSV 格式）") from exc

    headers, rows = build_rows(records, keys)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "发票信息"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F6FED")
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"

    # 列宽自适应 + 金额格式
    money_columns = {index for index, (key, _l) in enumerate(_columns(keys)) if key in config.MONEY_FIELDS}
    for index, header in enumerate(headers, start=1):
        width = max(10, min(46, len(str(header)) * 2 + 6))
        for row in rows[:200]:
            width = max(width, min(46, len(str(row[index - 1])) + 4))
        sheet.column_dimensions[get_column_letter(index)].width = width
        if index - 1 in money_columns:
            for row_index in range(2, len(rows) + 2):
                sheet.cell(row=row_index, column=index).number_format = "#,##0.00"

    if with_summary and records:
        sums = amount_sums(records)
        sheet.append([])
        line = _summary_line(headers, sums)
        sheet.append(line)
        last_row = sheet.max_row
        for column in range(1, len(headers) + 1):
            sheet.cell(row=last_row, column=column).font = Font(bold=True)
    workbook.save(path)
    return path


def export_json(
    path: str,
    records: Sequence[InvoiceRecord],
    keys: Optional[Sequence[str]] = None,
    with_summary: bool = True,
) -> str:
    pairs = _columns(keys)
    data: Dict[str, Any] = {
        "app": config.APP_NAME,
        "version": config.APP_VERSION,
        "exported_at": _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "count": len(records),
    }
    if with_summary:
        data["totals"] = amount_sums(records)
    data["invoices"] = [
        {
            label: (getattr(record, key, None) if key in config.MONEY_FIELDS else record.value(key, formatted=False))
            for key, label in pairs
        }
        for record in records
    ]
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
    return path


# --------------------------------------------------------------------------
# 统一入口
# --------------------------------------------------------------------------
def export(
    path: str,
    records: Sequence[InvoiceRecord],
    keys: Optional[Sequence[str]] = None,
    *,
    encoding: str = "utf-8-sig",
    with_summary: bool = True,
) -> str:
    """按文件扩展名导出，返回实际写入的路径。"""
    if not records:
        raise ExportError("没有可导出的记录")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return export_csv(path, records, keys, encoding, with_summary)
    if ext in (".xlsx", ".xlsm"):
        return export_xlsx(path, records, keys, with_summary)
    if ext == ".json":
        return export_json(path, records, keys, with_summary)
    raise ExportError(f"不支持的导出格式：{ext or '（无扩展名）'}，请使用 .csv / .xlsx / .json")
