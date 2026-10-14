"""文本文件 / 表格台账读取。"""
from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple

import config
from core.textutil import normalize

# 文本文件编码候选
_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "latin-1")


def read_text_file(path: str) -> str:
    """按常见中文编码依次尝试读取文本文件。"""
    try:
        import charset_normalizer  # type: ignore

        with open(path, "rb") as handle:
            raw = handle.read()
        best = charset_normalizer.from_bytes(raw).best()
        if best is not None:
            return normalize(raw.decode(best.encoding or "utf-8", errors="replace"))
    except Exception:
        pass

    with open(path, "rb") as handle:
        raw = handle.read()
    for encoding in _ENCODINGS:
        try:
            return normalize(raw.decode(encoding))
        except UnicodeDecodeError:
            continue
    return normalize(raw.decode("utf-8", errors="replace"))


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("（", "(").replace("）", ")")
    text = text.replace(" ", "")
    return text


def map_header_row(header: List[Any]) -> Dict[int, str]:
    """把表头行映射为 {列号: 字段键}。"""
    mapping: Dict[int, str] = {}
    for index, cell in enumerate(header):
        text = _normalize_header(cell)
        if not text:
            continue
        for key, aliases in config.HEADER_ALIASES.items():
            if key in mapping.values():
                continue
            for alias in aliases:
                alias_norm = alias.replace(" ", "").lower()
                if alias_norm and (text == alias_norm or alias_norm in text):
                    mapping[index] = key
                    break
            if index in mapping:
                break
    return mapping


def has_header_columns(mapping: Dict[int, str]) -> bool:
    """至少识别出"发票号码"或"金额"相关列，才算结构化台账。"""
    values = set(mapping.values())
    return bool(values & {"number", "total", "amount"})


def read_table_file(path: str) -> Dict[str, Any]:
    """读取表格/CSV，返回 {"headers": [...], "rows": [...], "text": "...", "structured": bool}。

    - 若能识别出发票字段表头，则结构化返回，由调用方直接映射成记录；
    - 否则把所有单元格拼成文本，交给通用正则解析。
    """
    lower = path.lower()
    headers: List[str] = []
    rows: List[List[Any]] = []

    if lower.endswith((".xlsx", ".xlsm")):
        headers, rows = _read_xlsx(path)
    elif lower.endswith(".xls"):
        headers, rows = _read_xls(path)
    else:  # csv / txt：按分隔符猜
        headers, rows = _read_delimited(path)

    mapping = map_header_row(headers)
    structured = has_header_columns(mapping)

    text_parts: List[str] = []
    if headers:
        text_parts.append(" ".join(str(h) for h in headers if h not in (None, "")))
    for row in rows:
        cells = [str(cell) for cell in row if cell not in (None, "")]
        if cells:
            text_parts.append(" ".join(cells))

    return {
        "headers": headers,
        "rows": rows,
        "mapping": mapping,
        "structured": structured,
        "text": normalize("\n".join(text_parts)),
    }


def _read_xlsx(path: str) -> Tuple[List[str], List[List[Any]]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("读取 .xlsx 需要 openpyxl，请执行：pip install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows: List[List[Any]] = []
    for row in sheet.iter_rows(values_only=True):
        if row is None:
            continue
        if all(cell in (None, "") for cell in row):
            continue
        rows.append([cell for cell in row])
    workbook.close()
    if not rows:
        return [], []
    header = ["" if cell is None else str(cell) for cell in rows[0]]
    return header, rows[1:]


def _read_xls(path: str) -> Tuple[List[str], List[List[Any]]]:
    try:
        import xlrd  # type: ignore
    except Exception as exc:
        raise RuntimeError("读取 .xls 需要 xlrd，请执行：pip install xlrd") from exc

    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if not rows:
        return [], []
    return [str(cell) for cell in rows[0]], rows[1:]


def _read_delimited(path: str) -> Tuple[List[str], List[List[Any]]]:
    import csv
    import io

    content = read_text_file(path)
    sample = content[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(content), dialect)
    rows = [row for row in reader if any(str(cell).strip() for cell in row)]
    if not rows:
        return [], []
    return [str(cell) for cell in rows[0]], rows[1:]
